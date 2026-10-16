"""Extract private prompt presets and public cover images from the supplied workbook."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    workbook_path = project_root / '产品方案' / '0901_1015图片文字含源图_中英版.xlsx'
    cover_directory = project_root / 'public' / 'images' / 'image-templates'
    generated_catalog = (
        project_root
        / 'src'
        / 'lib'
        / 'image-creator'
        / 'spreadsheet-templates.server.ts'
    )

    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=True)
    worksheet = workbook['汇总']
    image_by_row = {
        image.anchor._from.row + 1: image
        for image in worksheet._images
        if image.anchor._from.col + 1 == 2
    }

    templates: list[dict[str, str]] = []
    cover_directory.mkdir(parents=True, exist_ok=True)
    for row in range(3, worksheet.max_row + 1):
        date = worksheet.cell(row, 1).value
        english_prompt = worksheet.cell(row, 7).value
        image = image_by_row.get(row)
        if (
            not isinstance(date, str)
            or not isinstance(english_prompt, str)
            or not english_prompt.strip()
            or not image
        ):
            continue

        image_name = f'{date}.png'
        (cover_directory / image_name).write_bytes(image._data())
        templates.append(
            {
                'id': f'image-template-{date}',
                'name': f'Editorial poster {date}',
                'description': (
                    f'A reference-led 3:4 editorial poster from the {date} collection.'
                ),
                'coverUrl': f'/images/image-templates/{image_name}',
                'basePrompt': english_prompt.strip(),
            }
        )

    templates.sort(key=lambda template: template['id'], reverse=True)
    if len(templates) != 37:
        raise RuntimeError(f'Expected 37 usable templates, found {len(templates)}.')

    lines = [
        "import type { ImageTemplate } from './templates.server';",
        '',
        '// Generated from 产品方案/0901_1015图片文字含源图_中英版.xlsx.',
        '// This is server-only: basePrompt is intentionally never returned to the browser.',
        'export const spreadsheetImageTemplateSeeds: readonly ImageTemplate[] = [',
    ]
    for template in templates:
        lines.extend(
            [
                '  {',
                f"    id: {json.dumps(template['id'], ensure_ascii=False)},",
                f"    name: {json.dumps(template['name'], ensure_ascii=False)},",
                "    category: 'Editorial',",
                "    tags: ['editorial', 'reference-photo', '3:4'],",
                f"    description: {json.dumps(template['description'], ensure_ascii=False)},",
                f"    coverUrl: {json.dumps(template['coverUrl'], ensure_ascii=False)},",
                f"    exampleUrl: {json.dumps(template['coverUrl'], ensure_ascii=False)},",
                "    inputPolicy: 'optional',",
                "    defaultMode: '2K_HD',",
                "    defaultRatio: '3:4',",
                '    defaultCount: 1,',
                '    revision: 1,',
                f"    basePrompt: {json.dumps(template['basePrompt'], ensure_ascii=False)},",
                "    negativePrompt: 'low quality, distorted details, watermark, unreadable text',",
                '  },',
            ]
        )
    lines.extend(['];', ''])
    generated_catalog.write_text('\n'.join(lines), encoding='utf-8')
    print(f'Extracted {len(templates)} templates.')


if __name__ == '__main__':
    main()
